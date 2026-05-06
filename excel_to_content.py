#!/usr/bin/env python3
"""
excel_to_content.py
Reads contents/member-info.xlsx and regenerates:
  - contents/structures/members.json
  - contents/structures/members/{webId}.json    (members who have page content)
  - contents/articles/members-about/{webId}.md
  - contents/articles/members-position/{webId}.md
  - contents/articles/members-interest/{webId}.md

WebID is the formula-derived key (lowercase, hyphens/spaces removed from Full Name).
It is used for all filenames, page URLs, and image paths — no student ID mapping needed.

Run: conda run -n E3website python excel_to_content.py
"""

import json, os, re
import openpyxl

SECTION_ORDER = ['Principal Investigator', 'Full Time', 'Ph.D. Students', 'Master Students', 'Alumni']

SECTION_META = {
    'Principal Investigator': {'form': 'L'},
    'Full Time':              {'form': 'M'},
    'Ph.D. Students':         {'form': 'M'},
    'Master Students':        {'form': 'M'},
    'Alumni': {
        'form': 'S',
        'filter': True,
        'filterId': 'lum',
        'filterTitle': 'Filter by Admission Year:',
        'filterDefaultCheck': 5,
    },
}

PI_EXTRA_LINKS = [
    {
        'icon': '/assets/sprite.svg#svg-location',
        'text': 'CERB 601',
        'link': 'https://maps.app.goo.gl/crYHNJhSwBzqt2VJ8',
    },
]

PI_DESCRIPTION = (
    'Department of Civil Engineering<br>'
    'Department of Chemical Engineering (Joint Appointment)<br>'
    'National Taiwan University'
)

# ── helpers ───────────────────────────────────────────────────────────────────

def display_name(full_name: str, nickname: str) -> str:
    """Return 'Full Name (Nickname)' if nickname adds info not already in the
    full name, otherwise return 'Full Name' as-is.

    Logic: normalise both to lowercase letters only, then skip the bracket if
    either string starts with the other (covers cases like nickname='Jianhern'
    vs full='Jian Hern Yeoh', or nickname='I-Yun Lisa Hsieh, PhD' vs full='I-Yun Lisa Hsieh').
    """
    if not nickname or not full_name:
        return full_name or nickname
    fn = re.sub(r'[^a-z]', '', full_name.lower())
    nk = re.sub(r'[^a-z]', '', nickname.lower())
    # Also skip bracket if the nickname (≥3 chars) appears anywhere inside the full name
    if fn.startswith(nk) or nk.startswith(fn) or (len(nk) >= 3 and nk in fn):
        return full_name
    return f'{full_name} ({nickname})'


def find_image(web_id: str) -> str:
    """Return /assets/… path for a member's image, or nobody.svg."""
    img_dir = 'contents/images/members'
    for fname in sorted(os.listdir(img_dir)):
        stem = os.path.splitext(fname)[0]
        if stem.lower() == web_id.lower() and not fname.startswith('.'):
            return f'/assets/members/{fname}'
    return '/assets/members/nobody.svg'


def get_icon(current_position: str, batch: str) -> str:
    pos = (current_position or '').strip()
    b   = (batch or '').strip()
    if pos == 'RA':
        return '/assets/sprite.svg#svg-assistant'
    if pos.startswith('PHD'):
        return '/assets/sprite.svg#svg-doctor-2' if b.startswith('D') else '/assets/sprite.svg#svg-doctor-1'
    if pos == 'MS 2':
        return '/assets/sprite.svg#svg-master-2'
    if pos in ('MS 1', 'MS 0', 'MS DD'):
        return '/assets/sprite.svg#svg-master-1'
    return '/assets/sprite.svg#svg-master-2'  # Grad / Alumni fallback


def adm_year_sort_key(entry: dict) -> int:
    y = entry.get('admissionYear', '')
    m = re.search(r'\d+', str(y))
    return int(m.group()) if m else 0


def write_text(path: str, content: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)


# ── read Excel (data_only to resolve WebID formula values) ───────────────────

wb = openpyxl.load_workbook('contents/member-info.xlsx', data_only=True)
ws = wb['Members']
headers = [c.value for c in ws[1]]

def col(row, name):
    idx = headers.index(name)
    val = row[idx]
    return val if val is not None else ''


def col_optional(row, name):
    """Return cell value for optional column. Returns None if column header is absent."""
    if name not in headers:
        return None
    return row[headers.index(name)]


# ── process rows ──────────────────────────────────────────────────────────────

members_by_section = {s: [] for s in SECTION_ORDER}
extra_alumni = []

json_written = []
md_written   = []

for row in ws.iter_rows(min_row=2, values_only=True):
    web_id = row[headers.index('WebID')]
    if not web_id:
        continue  # blank / divider row

    web_id    = str(web_id).strip()
    nickname  = str(col(row, 'Nickname')             or '').strip()
    chi_name  = str(col(row, 'Chinese Name')         or '').strip()
    full_name = str(col(row, 'Full Name')            or '').strip()
    email     = str(col(row, 'Preferred Email')      or '').strip()
    ntu_email = str(col(row, 'NTU Email')            or '').strip()
    section   = str(col(row, 'Website Section')      or '').strip()
    adm_year  = col(row, 'Admission Year')
    graduated = str(col(row, 'Graduated')            or '').strip()
    also_alumni = str(col(row, 'Also in Alumni')     or '').strip().upper() == 'TRUE'
    alumni_year = col(row, 'Alumni Admission Year')
    about     = str(col(row, 'About')                or '').strip()
    position  = str(col(row, 'Position / Education') or '').strip()
    interests = str(col(row, 'Research Interests')   or '').strip()
    curr_pos  = str(col(row, 'Current Position')     or '').strip()
    batch     = str(col(row, 'Batch')                or '').strip()
    meta_description = (str(col_optional(row, 'metaDescription') or '').strip()) or None

    has_page      = bool(about or position)
    display_email = email or ntu_email
    img_path      = find_image(web_id)
    icon          = get_icon(curr_pos, batch)
    is_pi         = (section == 'Principal Investigator')

    # ── per-member JSON ───────────────────────────────────────────────────────
    if has_page:
        if section == 'Principal Investigator':
            position_for_seo = 'Director of E3 Center, NTU'
        else:
            position_for_seo = position or None

        member_json = {
            'chiNameEng':      display_name(full_name, nickname),
            'pubName':         full_name,   # plain Full Name for publication matching
            'metaDescription': meta_description,
            'position':        position_for_seo,
        }
        if not is_pi:
            member_json['graduated'] = (graduated.lower() == 'true')

        links = []
        if display_email:
            links.append({
                'icon': '/assets/sprite.svg#svg-send',
                'text': display_email,
                'link': f'mailto:{display_email}',
            })
        if is_pi:
            links.extend(PI_EXTRA_LINKS)

        page_content = {'links': links}
        if position:
            page_content['positionSection'] = {
                'content': f'contents/articles/members-position/{web_id}.md'
            }
        if about:
            page_content['aboutSection'] = [
                {'sectionTitle': 'About',
                 'content': f'contents/articles/members-about/{web_id}.md'}
            ]
        if is_pi:
            page_content['PublicationSection'] = [
                {'sectionTitle': 'Journal Publications', 'publications': []}
            ]
        member_json['pageContent'] = page_content

        json_path = f'contents/structures/members/{web_id}.json'
        os.makedirs(os.path.dirname(json_path), exist_ok=True)
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(member_json, f, ensure_ascii=False, indent=4)
        json_written.append(json_path)

        if about:
            p = f'contents/articles/members-about/{web_id}.md'
            write_text(p, about)
            md_written.append(p)
        if position:
            p = f'contents/articles/members-position/{web_id}.md'
            write_text(p, position)
            md_written.append(p)
        if interests:
            p = f'contents/articles/members-interest/{web_id}.md'
            write_text(p, interests)
            md_written.append(p)

    # ── members.json entry ────────────────────────────────────────────────────
    if section not in members_by_section:
        print(f'  WARNING: unknown section "{section}" for {web_id}, skipping')
        continue

    if is_pi:
        entry = {
            'engName':     full_name,
            'chiName':     chi_name,
            'description': PI_DESCRIPTION,
            'imgPath':     img_path,
            'links':       [],
        }
        if display_email:
            entry['links'].append({
                'icon': '/assets/sprite.svg#svg-send',
                'text': display_email,
                'link': f'mailto:{display_email}',
            })
        entry['links'].extend(PI_EXTRA_LINKS)
        if has_page:
            entry['pageLink']      = f'/members/{web_id}'
            entry['pageStructure'] = f'contents/structures/members/{web_id}.json'
    else:
        entry = {
            'engName': nickname or full_name,
            'chiName': chi_name,
            'imgPath': img_path,
            'icon':    icon,
        }
        if adm_year:
            entry['admissionYear'] = str(adm_year)
        if has_page:
            entry['pageLink']      = f'/members/{web_id}'
            entry['pageStructure'] = f'contents/structures/members/{web_id}.json'

    members_by_section[section].append(entry)

    if also_alumni and section != 'Alumni':
        alumni_entry = {
            'engName': nickname or full_name,
            'chiName': chi_name,
            'imgPath': img_path,
            'icon':    '/assets/sprite.svg#svg-master-2',
        }
        if alumni_year:
            alumni_entry['admissionYear'] = str(alumni_year)
        if has_page:
            alumni_entry['pageLink']      = f'/members/{web_id}'
            alumni_entry['pageStructure'] = f'contents/structures/members/{web_id}.json'
        extra_alumni.append(alumni_entry)

members_by_section['Alumni'].extend(extra_alumni)
members_by_section['Alumni'].sort(key=adm_year_sort_key)

# ── write members.json ────────────────────────────────────────────────────────

output = []
for section_title in SECTION_ORDER:
    members_list = members_by_section[section_title]
    if not members_list:
        continue
    section_obj = {'sectionTitle': section_title}
    section_obj.update(SECTION_META[section_title])
    section_obj['members'] = members_list
    output.append(section_obj)

members_json_path = 'contents/structures/members.json'
with open(members_json_path, 'w', encoding='utf-8') as f:
    json.dump(output, f, ensure_ascii=False, indent=4)

# ── summary ───────────────────────────────────────────────────────────────────

print(f'Written: {members_json_path}')
print(f'Written: {len(json_written)} member JSON files')
print(f'Written: {len(md_written)} markdown files')
print()
for s in SECTION_ORDER:
    print(f'  {s}: {len(members_by_section[s])} members')
