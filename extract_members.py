"""
extract_members.py
Extracts all member data (roster + page content + markdown) into a single Excel file.
Run: python extract_members.py
Output: contents/member-info-full.xlsx
"""

import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── helpers ──────────────────────────────────────────────────────────────────

def read_md(path):
    if path and os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return f.read().strip()
    return ''

def load_json(path):
    if path and os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

# ── load data ─────────────────────────────────────────────────────────────────

with open('contents/structures/members.json', 'r', encoding='utf-8') as f:
    sections = json.load(f)

# ── build rows ────────────────────────────────────────────────────────────────

rows = []

for section in sections:
    section_title = section.get('sectionTitle', '')
    for member in section.get('members', []):
        # Basic roster fields
        eng_name   = member.get('engName', '')
        chi_name   = member.get('chiName', '')
        img_path   = member.get('imgPath', '')
        page_link  = member.get('pageLink', '')
        page_struct = member.get('pageStructure', '')
        admission  = member.get('admissionYear', '')

        # Derive ID from pageLink or pageStructure
        member_id = ''
        if page_link:
            member_id = page_link.rstrip('/').split('/')[-1]
        elif page_struct:
            member_id = os.path.splitext(os.path.basename(page_struct))[0]

        # Load individual page JSON if it exists
        details = load_json(page_struct) if page_struct else {}
        pub_name      = details.get('pubName', '')
        chi_name_eng  = details.get('chiNameEng', '')
        graduated     = details.get('graduated', '')
        page_content  = details.get('pageContent', {})

        # Email — first link with svg-send icon
        email = ''
        for link in page_content.get('links', []):
            if 'svg-send' in link.get('icon', ''):
                email = link.get('text', '')
                break

        # Markdown content
        about_md    = ''
        pos_section = page_content.get('positionSection', {})
        pos_md_path = pos_section.get('content', '') if isinstance(pos_section, dict) else ''
        position_md = read_md(pos_md_path)

        for about_sec in page_content.get('aboutSection', []):
            about_md_path = about_sec.get('content', '')
            about_md = read_md(about_md_path)
            break  # only first about section

        interest_md = ''
        if member_id:
            interest_md = read_md(f'contents/articles/members-interest/{member_id}.md')

        rows.append({
            'Section':       section_title,
            'Nickname':      eng_name,
            'Chinese Name':  chi_name,
            'Full Name':     chi_name_eng,
            'pub Name':      pub_name,
            'ID':            member_id,
            'Email':         email,
            'Admission Year': admission,
            'Graduated':     str(graduated) if graduated != '' else '',
            'About':         about_md,
            'Position / Education': position_md,
            'Research Interests':   interest_md,
        })

# ── build Excel ───────────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Members'

HEADERS = [
    'Section', 'Nickname', 'Chinese Name', 'Full Name', 'pub Name',
    'ID', 'Email', 'Admission Year', 'Graduated',
    'About', 'Position / Education', 'Research Interests',
]

# Section colour map
SECTION_COLORS = {
    'Principal Investigator': 'D6E4F7',
    'Full Time':              'D6F7E4',
    'Ph.D. Students':         'FFF3CC',
    'Master Students':        'FCE4D6',
    'Alumni':                 'EDEDED',
}

header_fill = PatternFill('solid', fgColor='2F5496')
header_font = Font(bold=True, color='FFFFFF', size=11)
wrap        = Alignment(wrap_text=True, vertical='top')
thin        = Side(style='thin', color='CCCCCC')
border      = Border(left=thin, right=thin, top=thin, bottom=thin)

# Write headers
for col, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill   = header_fill
    cell.font   = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border

ws.row_dimensions[1].height = 20

# Write data rows
for row_idx, data in enumerate(rows, 2):
    section = data['Section']
    row_fill = PatternFill('solid', fgColor=SECTION_COLORS.get(section, 'FFFFFF'))

    for col, key in enumerate(HEADERS, 1):
        cell = ws.cell(row=row_idx, column=col, value=data.get(key, ''))
        cell.fill      = row_fill
        cell.alignment = wrap
        cell.border    = border

    # Taller rows for markdown content columns
    has_content = any(data.get(k) for k in ('About', 'Position / Education', 'Research Interests'))
    ws.row_dimensions[row_idx].height = 80 if has_content else 18

# Column widths
col_widths = {
    'Section': 18, 'Nickname': 14, 'Chinese Name': 14, 'Full Name': 22,
    'pub Name': 20, 'ID': 16, 'Email': 28, 'Admission Year': 14,
    'Graduated': 10, 'About': 50, 'Position / Education': 50,
    'Research Interests': 50,
}
for col, key in enumerate(HEADERS, 1):
    ws.column_dimensions[get_column_letter(col)].width = col_widths.get(key, 15)

# Freeze header row
ws.freeze_panes = 'A2'

# ── save ──────────────────────────────────────────────────────────────────────

out_path = 'contents/member-info-full.xlsx'
wb.save(out_path)
print(f"Saved {len(rows)} members → {out_path}")
