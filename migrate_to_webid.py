#!/usr/bin/env python3
"""
migrate_to_webid.py  (one-time migration script)
- Renames contents/images/members/* from old names (student IDs / old WebIDs) to new WebIDs
- Deletes old student-ID-named JSON and MD files so excel_to_content.py rebuilds cleanly

Run once: conda run -n E3website python migrate_to_webid.py
"""

import os, json, shutil
import openpyxl

IMG_DIR      = 'contents/images/members'
MEMBERS_JSON = 'contents/structures/members.json'
EXCEL_FILE   = 'contents/member-info.xlsx'

MD_DIRS = [
    'contents/articles/members-about',
    'contents/articles/members-position',
    'contents/articles/members-interest',
]
JSON_DIR = 'contents/structures/members'

# ── 1. Build chi_name → new_webid from new Excel ─────────────────────────────
wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb['Members']
headers = [c.value for c in ws[1]]
chi_idx = headers.index('Chinese Name')
web_idx = headers.index('WebID')

new_webid_by_chi = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    chi, wid = row[chi_idx], row[web_idx]
    if chi and wid:
        new_webid_by_chi[chi] = str(wid).strip()

# ── 2. Build chi_name → old_img_stem from members.json ──────────────────────
with open(MEMBERS_JSON) as f:
    sections = json.load(f)

old_img_stem_by_chi = {}  # chi_name → stem of imgPath filename
old_file_id_by_chi  = {}  # chi_name → pageStructure stem (for JSON/MD cleanup)

for s in sections:
    for m in s.get('members', []):
        chi     = m.get('chiName', '')
        img     = m.get('imgPath', '')
        ps      = m.get('pageStructure', '')
        if chi and img:
            stem = os.path.splitext(os.path.basename(img))[0]
            old_img_stem_by_chi[chi] = stem
        if chi and ps:
            old_file_id_by_chi[chi] = os.path.splitext(os.path.basename(ps))[0]

# ── 3. Build image rename plan: old_stem → new_webid ────────────────────────
# Build extension lookup from actual files
ext_by_stem = {}
for fname in os.listdir(IMG_DIR):
    if fname.startswith('.'): continue
    stem, ext = os.path.splitext(fname)
    ext_by_stem[stem.lower()] = (stem, ext, fname)  # preserve original casing

rename_plan  = []   # (old_fname, new_fname)
no_image     = []   # chi names with no image found
no_new_webid = []   # chi names with no new WebID in Excel

for chi, old_stem in old_img_stem_by_chi.items():
    new_wid = new_webid_by_chi.get(chi)
    if not new_wid:
        no_new_webid.append(chi)
        continue
    if old_stem == new_wid:
        continue  # already correct
    match = ext_by_stem.get(old_stem.lower())
    if not match:
        no_image.append((chi, old_stem))
        continue
    _, ext, old_fname = match
    new_fname = new_wid + ext
    rename_plan.append((old_fname, new_fname))

# ── 4. Print plan before executing ───────────────────────────────────────────
print('=== Image renames ===')
for old, new in sorted(rename_plan):
    print(f'  {old} → {new}')

if no_image:
    print('\nWARNING: image file not found for:')
    for chi, stem in no_image:
        print(f'  {chi} (expected stem: {stem})')

if no_new_webid:
    print('\nNOTE: no new WebID in Excel for (image kept as-is):')
    for chi in no_new_webid:
        print(f'  {chi}')

# ── 5. Execute image renames ─────────────────────────────────────────────────
print('\nRenaming images...')
for old_fname, new_fname in rename_plan:
    src = os.path.join(IMG_DIR, old_fname)
    dst = os.path.join(IMG_DIR, new_fname)
    os.rename(src, dst)
    print(f'  ✓ {old_fname} → {new_fname}')

# ── 6. Delete old JSON and MD files (will be regenerated with new names) ─────
print('\nDeleting old JSON files...')
for chi, old_id in old_file_id_by_chi.items():
    new_wid = new_webid_by_chi.get(chi)
    if new_wid and old_id != new_wid:
        path = os.path.join(JSON_DIR, old_id + '.json')
        if os.path.exists(path):
            os.remove(path)
            print(f'  ✓ deleted {path}')

print('\nDeleting old MD files...')
for chi, old_id in old_file_id_by_chi.items():
    new_wid = new_webid_by_chi.get(chi)
    if new_wid and old_id != new_wid:
        for md_dir in MD_DIRS:
            path = os.path.join(md_dir, old_id + '.md')
            if os.path.exists(path):
                os.remove(path)
                print(f'  ✓ deleted {path}')

print('\nMigration complete. Run build.py to regenerate content.')
