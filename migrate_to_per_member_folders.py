"""One-time migration: extract existing per-member content from the legacy
22-column Excel + members-{about,position,interest}/*.md into the new
contents/members/{webId}/ folder structure.

Run from the repo root:
    python migrate_to_per_member_folders.py
    python migrate_to_per_member_folders.py --force      # overwrite existing folders
    python migrate_to_per_member_folders.py --dry-run

Idempotent: re-running without --force will refuse to clobber existing
content folders.
"""
from __future__ import annotations


import json as _json
import shutil
from pathlib import Path
from typing import Mapping, Any


def parse_interests_from_slash_md(md: str) -> list[str]:
    """Parse the existing /Topic A\\n/Topic B format into ['Topic A', 'Topic B']."""
    out: list[str] = []
    for line in (md or "").splitlines():
        s = line.strip()
        if s.startswith("/"):
            out.append(s.lstrip("/").strip())
    return out


def _s(value: Any) -> str:
    """Coerce Excel cell value → trimmed string, empty when None/blank."""
    if value is None:
        return ""
    return str(value).strip()


def build_member_json(row: Mapping[str, Any]) -> dict[str, Any]:
    """Build the all-keys-present member.json dict from a legacy Excel row.

    `row` is keyed by the legacy column names (e.g. 'NTU Email',
    'Google Scholar'). Missing columns default to empty values — never raises.
    Note: the legacy Excel has no Office/Office Map columns; only the PI has
    an office link, special-cased by main()."""
    return {
        "position": _s(row.get("Position / Education")),
        "email": {
            "ntu":       _s(row.get("NTU Email")),
            "preferred": _s(row.get("Preferred Email")),
        },
        "interests": parse_interests_from_slash_md(_s(row.get("Research Interests"))),
        "links": {
            "scholar":      _s(row.get("Google Scholar")),
            "orcid":        _s(row.get("ORCID")),
            "linkedin":     _s(row.get("LinkedIn")),
            "researchgate": _s(row.get("ResearchGate")),
            "ntu_scholars": _s(row.get("NTU Scholars")),
            "facebook":     _s(row.get("Facebook")),
            "office": {
                "text": _s(row.get("Office")),
                "url":  _s(row.get("Office Map")),
            },
        },
        "metaDescription": _s(row.get("metaDescription")),
    }


def write_member_folder(
    *,
    dest: Path,
    member_json: dict,
    about_md_src: Path | None,
    photo_src: Path | None,
) -> None:
    """Create dest/ and write member.json (+ optionally about.md, photo.{ext}).

    Files are overwritten if they already exist (caller controls --force at
    the CLI level; at the unit level this function is unconditional)."""
    dest.mkdir(parents=True, exist_ok=True)

    (dest / "member.json").write_text(
        _json.dumps(member_json, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    if about_md_src is not None and about_md_src.exists():
        (dest / "about.md").write_text(
            about_md_src.read_text(encoding="utf-8"),
            encoding="utf-8",
        )

    if photo_src is not None and photo_src.exists():
        ext = photo_src.suffix.lower()
        shutil.copyfile(photo_src, dest / f"photo{ext}")


ADMIN_COLUMNS: list[str] = [
    "WebID",
    "Full Name",
    "Nickname",
    "Chinese Name",
    "Website Section",
    "Admission Year",
    "Graduated",
    "Also in Alumni",
    "Alumni Admission Year",
    "Current Position",
    "Batch",
]


def write_slim_excel(legacy_path: Path, slim_path: Path) -> None:
    """Read the legacy workbook; write a new workbook keeping only the 11
    ADMIN_COLUMNS (in the order defined by that constant)."""
    import openpyxl
    src_wb = openpyxl.load_workbook(legacy_path, data_only=True)
    src_ws = src_wb["Members"]
    src_headers = [c.value for c in src_ws[1]]

    keep_indices = []
    for col in ADMIN_COLUMNS:
        if col not in src_headers:
            raise ValueError(f"Legacy Excel missing required admin column: {col!r}")
        keep_indices.append(src_headers.index(col))

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "Members"
    out_ws.append(ADMIN_COLUMNS)
    for row in src_ws.iter_rows(min_row=2, values_only=True):
        out_ws.append([row[i] for i in keep_indices])

    out_wb.save(slim_path)


_EMPTY_MEMBER_JSON: dict = {
    "position": "",
    "email": {"ntu": "", "preferred": ""},
    "interests": [],
    "links": {
        "scholar": "",
        "orcid": "",
        "linkedin": "",
        "researchgate": "",
        "ntu_scholars": "",
        "facebook": "",
        "office": {"text": "", "url": ""},
    },
    "metaDescription": "",
}

_TEMPLATE_ABOUT_MD = """# Your name here

Write your bio in this file. You can use Markdown:

- **bold** with `**bold**`
- *italics* with `*italics*`
- [links](https://example.com) with `[text](url)`

Multiple paragraphs are fine. This is the prose that appears in the "About"
section of your member page.

(Delete this placeholder text and write your bio here.)
"""

_TEMPLATE_README_MD = """# Updating your E3 Center profile

Edit **two files** and send them back to the maintainer:

1. **`member.json`** — your position, emails, interests, profile links.
   - Replace the string values. Don't delete keys or rearrange.
   - Leave any field as `""` to hide it on the website.
   - `interests` is a list — add or remove items as you like (1–5 typical).
2. **`about.md`** — your biographical paragraphs. Plain Markdown:
   - `**bold**`, `*italics*`, `[links](https://example.com)`
   - Blank line = new paragraph
   - `- item` for bullets

Send both files back when done. Do not edit anything in `member-info.xlsx`
— your admin info (name, batch, section) is handled separately.
"""


def write_member_template(dest: Path) -> None:
    """Create dest/ (typically contents/MEMBER_TEMPLATE/) with the three
    skeleton files maintainers send to new members."""
    dest.mkdir(parents=True, exist_ok=True)
    (dest / "member.json").write_text(
        _json.dumps(_EMPTY_MEMBER_JSON, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    (dest / "about.md").write_text(_TEMPLATE_ABOUT_MD, encoding="utf-8")
    (dest / "README.md").write_text(_TEMPLATE_README_MD, encoding="utf-8")


import argparse
import sys

# The PI's office link is hardcoded in excel_to_content.py's PI_EXTRA_LINKS,
# not an Excel column. Mirror it here so the PI's member.json keeps it.
_PI_OFFICE = {"text": "CERB 601", "url": "https://maps.app.goo.gl/crYHNJhSwBzqt2VJ8"}


def _find_legacy_photo(image_dir: Path, web_id: str) -> Path | None:
    """Find contents/images/members/{webId}.{ext} case-insensitively
    (handles shaoyangcheung.JPG etc.). Returns the first match or None."""
    if not image_dir.is_dir():
        return None
    for p in sorted(image_dir.iterdir()):
        if p.is_file() and not p.name.startswith(".") and p.stem.lower() == web_id.lower():
            return p
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="One-time migration to per-member folders.")
    parser.add_argument("--force", action="store_true",
                        help="Overwrite existing contents/members/*/ folders.")
    parser.add_argument("--dry-run", action="store_true",
                        help="Print what would happen without writing files.")
    args = parser.parse_args(argv)

    repo_root     = Path(".")
    legacy_xlsx   = repo_root / "contents" / "member-info.xlsx"
    archive_xlsx  = repo_root / "contents" / "member-info.legacy.xlsx"
    members_root  = repo_root / "contents" / "members"
    template_root = repo_root / "contents" / "MEMBER_TEMPLATE"
    old_about_dir = repo_root / "contents" / "articles" / "members-about"
    old_image_dir = repo_root / "contents" / "images" / "members"

    if not legacy_xlsx.exists():
        print(f"✗ {legacy_xlsx} does not exist", file=sys.stderr)
        return 2

    if members_root.exists() and any(members_root.iterdir()) and not args.force:
        print(f"✗ {members_root}/ is not empty. Re-run with --force to overwrite.",
              file=sys.stderr)
        return 2

    import openpyxl
    wb = openpyxl.load_workbook(legacy_xlsx, data_only=True)
    ws = wb["Members"]
    headers = [c.value for c in ws[1]]

    n_processed = 0
    for row_tuple in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, row_tuple))
        web_id = _s(row.get("WebID"))
        if not web_id:
            continue  # blank / divider row

        member_data = build_member_json(row)

        # PI office link is hardcoded upstream, not an Excel column.
        if _s(row.get("Website Section")) == "Principal Investigator":
            member_data["links"]["office"] = dict(_PI_OFFICE)

        about_src = old_about_dir / f"{web_id}.md"
        about_src = about_src if about_src.exists() else None

        photo_src = _find_legacy_photo(old_image_dir, web_id)

        dest = members_root / web_id

        if args.dry_run:
            print(f"  would write {dest}/  "
                  f"(about={about_src is not None}, photo={photo_src is not None})")
        else:
            write_member_folder(
                dest=dest,
                member_json=member_data,
                about_md_src=about_src,
                photo_src=photo_src,
            )
            print(f"  ✓ {dest}/")
        n_processed += 1

    if args.dry_run:
        print(f"\n(dry-run) would process {n_processed} members.")
        return 0

    # Slim Excel + archive original
    tmp_slim = legacy_xlsx.with_suffix(".slim.xlsx")
    write_slim_excel(legacy_xlsx, tmp_slim)
    legacy_xlsx.rename(archive_xlsx)
    tmp_slim.rename(legacy_xlsx)
    print(f"  ✓ wrote slim {legacy_xlsx}")
    print(f"  ✓ archived original → {archive_xlsx}")

    # Template skeleton
    write_member_template(template_root)
    print(f"  ✓ wrote {template_root}/")

    print(f"\nDone. Processed {n_processed} members.")
    print("\nNext steps (manual):")
    print("  - Inspect a few migrated folders (e.g. contents/members/iyunlisahsieh/)")
    print("  - Run: python build.py  (should succeed with at most warnings)")
    print("  - diff -r /tmp/docs.before docs   (zero diffs expected)")
    print("  - git rm -r the now-stale legacy member directories")
    return 0


if __name__ == "__main__":
    sys.exit(main())
